from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import pdfplumber
from docx import Document
from pypdf import PdfReader

from .common import check_docx, finding
from .evidence import block_text


def inspect_docx(path, draft):
    check_docx(path)
    doc = Document(path)
    issues = []
    texts = [p.text for p in doc.paragraphs]
    texts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    texts += [p.text for s in doc.sections for container in (s.header, s.footer) for p in container.paragraphs]
    text = "\n".join(texts)
    if not text.strip():
        issues.append(finding("DOCX_EMPTY", "Document text empty"))
    if draft["title"] not in texts:
        issues.append(finding("DOCX_TITLE", "Document title missing"))
    if re.search(r"TODO|XXX|\{\{|\{%|待补充", text):
        issues.append(finding("DOCX_PLACEHOLDER", "Unresolved placeholder in final document"))
    for section in draft["sections"]:
        if section["title"] not in texts:
            issues.append(finding("DOCX_SECTION", f"Heading missing: {section['title']}"))
        for block in section["blocks"]:
            required = [block["text"]] if block["type"] == "paragraph" else [str(c) for row in [block["headers"]] + block["rows"] for c in row]
            if any(t not in texts for t in required):
                issues.append(finding("DOCX_CONTENT", f"Content missing: {block['id']}"))
    available = min(s.page_width - s.left_margin - s.right_margin for s in doc.sections)
    for i, table in enumerate(doc.tables, 1):
        if sum(col.width or 0 for col in table.columns) > available + 12700:
            issues.append(finding("TABLE_OVERFLOW", f"Table {i} grid wider than page content"))
    return issues


def preview(docx, folder, soffice=None, pdftoppm=None):
    """Run only supplied/local executables; no download or cloud fallback."""
    soffice = soffice or shutil.which("soffice") or shutil.which("soffice.exe")
    pdftoppm = pdftoppm or shutil.which("pdftoppm")
    if not soffice:
        raise RuntimeError("LibreOffice unavailable: pass --soffice PATH")
    if not pdftoppm:
        raise RuntimeError("Poppler unavailable: pass --pdftoppm PATH")
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="word-skills-lo-") as temp:
        work = Path(temp)
        profile = (work / "profile").as_uri()
        command = [str(soffice), f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf", "--outdir", str(work), str(Path(docx).resolve())]
        result = subprocess.run(command, capture_output=True, timeout=90, text=True, encoding="utf-8", errors="replace")
        pdf = work / (Path(docx).stem + ".pdf")
        if result.returncode or not pdf.is_file() or pdf.stat().st_size == 0:
            raise RuntimeError(f"LibreOffice conversion failed: {result.stdout} {result.stderr}")
        destination = folder / "preview.pdf"
        shutil.copyfile(pdf, destination)
    result = subprocess.run([str(pdftoppm), "-png", "-r", "110", str(destination), str(folder / "page")], capture_output=True, timeout=90)
    images = sorted(folder.glob("page-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
    if result.returncode or len(images) != len(PdfReader(destination).pages):
        raise RuntimeError("PDF page image rendering failed or incomplete")
    version = subprocess.run([str(soffice), "--version"], capture_output=True, text=True, timeout=20).stdout.strip()
    return {"backend": "LibreOffice", "version": version, "pages": len(images), "page_images": [p.name for p in images]}


def inspect_pdf(path, draft):
    issues, pages, fonts = [], [], set()
    headings = {s["title"].replace(" ", "") for s in draft["sections"]}
    with pdfplumber.open(path) as pdf:
        if not pdf.pages:
            issues.append(finding("PDF_EMPTY", "No PDF pages"))
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            pages.append({"page": i, "characters": len(text), "width": page.width, "height": page.height})
            if len(re.sub(r"[\s\d]", "", text)) < 3:
                issues.append(finding("BLANK_PAGE", f"Page {i} blank or nearly empty"))
            if "\ufffd" in text or "(cid:" in text:
                issues.append(finding("PDF_GLYPH", f"Page {i} contains replacement glyph markers"))
            lines = [s.strip().replace(" ", "") for s in text.splitlines() if s.strip() and not s.strip().isdigit()]
            if lines and lines[-1] in headings:
                issues.append(finding("ORPHAN_HEADING", f"Heading at bottom of page {i}"))
            for char in page.chars:
                fonts.add(char.get("fontname", "unknown"))
                if char["x0"] < -1 or char["x1"] > page.width + 1 or char["top"] < -1 or char["bottom"] > page.height + 1:
                    issues.append(finding("PDF_BOUNDS", f"Text outside page {i}"))
                    break
    return {"issues": issues, "pages": pages, "rendered_fonts": sorted(fonts),
            "font_substitution": "not automatically determined; compare rendered fonts and page images with template",
            "limits": "Page bounds and blank/orphan heuristics; overlaps, cell clipping and meaning require visual/content review"}
