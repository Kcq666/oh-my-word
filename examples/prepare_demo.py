"""Prepare two reviewed-source demo drafts from the user's existing patent project.

This is a fixed case fixture, not an automatic fact verifier or general writer.
Statements below were selected against the actual source during this implementation.
It refuses different source text instead of manufacturing replacement evidence.
"""
from pathlib import Path
import argparse

from word_skills.common import now, read_json, write_json
from word_skills.evidence import contract
from word_skills.ingest import ingest
from word_skills.render import default_document
from docx.shared import Mm, Pt


def prepare(project, output):
    project, output = Path(project).resolve(), Path(output).resolve()
    if output.exists():
        raise ValueError("Use a new demo path")
    paths = ["pyproject.toml", "oh_my_patent/ingest.py", "oh_my_patent/render/docx.py", "oh_my_patent/gate.py", "oh_my_patent/oxml.py"]
    case = output / "case"
    ingest([project / p for p in paths], case)
    source_by_name = {Path(s["original_path"]).name: s for s in read_json(case / "source-index.json")["sources"]}
    claims = []

    def claim(cid, name, start, count, statement):
        source = source_by_name[name]
        segments = read_json(case / source["segments"])
        matches = [i for i, s in enumerate(segments) if start in s["text"]]
        if len(matches) != 1:
            raise ValueError(f"Source changed or ambiguous quote: {name}: {start}")
        selected = segments[matches[0]:matches[0] + count]
        claims.append({"id": cid, "statement": statement, "source_id": source["id"], "locators": [s["locator"] for s in selected],
                       "quote": "\n".join(s["text"] for s in selected), "status": "verified",
                       "verified_by": "Codex reviewed fixed source fixture", "verified_at": now()})

    claim("C1", "pyproject.toml", 'version = "0.2.0"', 1, "项目声明版本为 0.2.0。")
    claim("C2", "pyproject.toml", 'requires-python = ">=3.10"', 1, "项目要求 Python 3.10 或更高版本。")
    claim("C3", "ingest.py", '"""把项目材料（Word / PowerPoint）', 1, "材料摄入模块将 Word 和 PowerPoint 资料转换为 Markdown，供后续处理。")
    claim("C4", "ingest.py", "- 走 mammoth 的 **HTML** 输出", 2, "Word 摄入先使用 mammoth 生成 HTML，再交给 markdownify 转成 Markdown，以保留表格结构。")
    claim("C5", "ingest.py", "- 公式的**文本近似是退化的**", 2, "公式提取只能得到退化的文本近似，每处公式都产生 Notice，需核对原式。")
    claim("C6", "docx.py", "def render_docx(", 7, "render_docx 接受 PatentDraft、输出路径、可选 RenderOptions 和 split 开关，返回文件路径列表。")
    claim("C7", "gate.py", "if not facts_path or not review_path:", 3, "最终审阅交付要求提供技术事实账本和内容审查文件，缺少时返回错误。")
    claim("C8", "gate.py", 'if review.get("draft_sha256") != sha256(draft_path):', 4, "稿件哈希或证据摘要变化时，原内容审查被判定为过期。")
    claim("C9", "gate.py", 'fail("FINAL-MATH",', 1, "当前版本不提供复杂公式排版，原始 LaTeX 等内容需要先恢复为受支持的表示。")
    claim("C10", "oxml.py", "section.page_width = Mm(210)", 6, "专利页面宽 210 mm、高 297 mm，上边距 25 mm、左边距 25 mm、右边距 15 mm、下边距 15 mm。")
    claim("C11", "gate.py", "# final 不可由 --force 绕过", 10, "最终检查逐幅读取图件；未登记路径或图件不可解码时拒绝通过。")
    write_json(case / "evidence.json", {"schema_version": 1, "claims": claims})
    indexed = {c["id"]: c["statement"] for c in claims}

    def fact(bid, *ids):
        return {"id": bid, "type": "paragraph", "intent": "fact", "text": "".join(indexed[i] for i in ids), "claim_ids": list(ids)}

    def note(bid, text, intent="analysis"):
        return {"id": bid, "type": "paragraph", "intent": intent, "text": text, "claim_ids": []}

    report_blocks = {
        "background": [note("r-background", "本报告面向维护专利文档工具的工程人员，说明材料摄入、Word 输出和审阅交付检查的实现边界。结论是现有工具已形成相互衔接的处理模块，通用技术文档可以复用部分基础能力，但应保留专利版式与业务规则的独立性。"), fact("r-version", "C1", "C3")],
        "objectives": [note("r-goals", "本次分析关注材料转换是否显式暴露信息损失、文档渲染是否具有清楚的输入输出，以及修改稿件后能否避免复用过期审查。它服务于后续文档能力扩展，不用于判断任何具体专利的授权前景。"), fact("r-review", "C7")],
        "methods": [note("r-method", "分析直接阅读项目配置、材料摄入模块、渲染入口、最终检查函数和页面配置函数。配置声明用于解释运行前提，代码分支用于解释检查逻辑；二者均不能替代目标环境中的实际压力测试或用户验收。"), fact("r-environment", "C2", "C6")],
        "implementation": [fact("r-ingest", "C4", "C5"), fact("r-hash", "C8", "C11"), {"id": "r-table", "type": "table", "intent": "fact", "headers": ["页面参数", "当前设置"], "rows": [["纸张宽高", "210 mm × 297 mm"], ["上边距与左边距", "25 mm"], ["右边距与下边距", "15 mm"]], "widths": [2, 3], "claim_ids": ["C10"]}],
        "results": [note("r-observed", "源码层面的观察表明，转换告警、图片读取和审查摘要检测都有明确的处理位置。该观察支持对现有实现进行职责划分；它不证明所有输入都能无损转换，也不说明生成文件在每种办公软件中具有相同分页。"), fact("r-limit", "C9"), note("r-gap", "本次材料没有提供吞吐量、延迟、并发容量或生成质量的实测数据，因此本报告不作性能达标结论。", "gap")],
        "risks": [note("r-risk", "建议先提取与业务无关的来源登记和字体配置，再单独建立通用文档的页面配置。公式、图片和复杂模板应保留原件核对路径；所有内容结论仍需结合实际来源审阅。扩展过程中应通过专利回归样例确认原格式未受影响。"), note("r-conclusion", "现有项目适合作为文档工具体系的工程基础。可复用的重点是可追溯输入和显式检查机制，专利特有的章节、权项支持关系及页面参数应继续由专利模块管理。")]
    }
    design_blocks = {
        "requirements": [note("d-scope", "本说明记录现有专利文档工具的主要结构，供维护和扩展时使用。目标是让材料读取、文档输出与审查证据各有明确职责；未来的通用报告生成属于扩展设计，不应被描述为现有专利工具已经具备的业务功能。"), fact("d-version", "C1")],
        "constraints": [fact("d-runtime", "C2", "C10"), note("d-constraint", "专利页面参数是现有格式约束。扩展为技术报告时应引入独立版式配置，而不直接修改专利默认值。所有未经目标办公软件验证的兼容性结论都应保持有限范围。")],
        "architecture": [fact("d-flow", "C3", "C4", "C6"), {"id": "d-table", "type": "table", "intent": "fact", "headers": ["模块", "职责与边界"], "rows": [["材料摄入", "Word 与 PowerPoint 转为 Markdown，公式退化产生告警。"], ["文档渲染", "根据 PatentDraft 和 RenderOptions 写出 Word 路径列表。"], ["最终检查", "校验事实账本、内容审查、图件及摘要变化。"]], "widths": [1, 4], "claim_ids": ["C3", "C5", "C6", "C7", "C8", "C11"]}],
        "interfaces": [fact("d-api", "C6", "C7"), note("d-contract", "渲染入口的输出是文件路径列表，调用方应以实际生成的文件为后续检查对象。审查账本属于独立输入，不应通过渲染成功反推内容已经审阅。接口扩展宜保持旧入口的调用含义和错误语义。")],
        "deployment": [fact("d-failure", "C5", "C8", "C11"), note("d-recovery", "发生材料损失或审查过期时，恢复路径应回到原件核对和相关内容重审。单纯更新哈希不能修复已经失效的语义结论。建议在新的输出目录生成修订版，以便保留此前稿件和对应审查记录。")],
        "acceptance": [note("d-accept", "验收应覆盖材料告警是否保留、输出文件是否可打开、缺少图件和审查材料是否失败，以及修改稿件后旧审查是否失效。版面验收应实际渲染文档并检查每页，而不是只判断 XML 标签或文件存在。"), fact("d-math", "C9"), note("d-open", "复杂公式、跨软件分页及特定客户模板仍需要有代表性的输入和目标环境验证。扩展通用文档能力时，继续以现有专利样例作为回归依据。", "gap")]
    }
    for mode, title, blocks in [("technical-report", "专利文档工具技术报告", report_blocks), ("design-spec", "专利文档工具设计说明", design_blocks)]:
        schema = contract(mode)
        draft = {"schema_version": 1, "mode": mode, "title": title, "sections": [{"id": k, "title": v, "blocks": blocks[k]} for k, v in schema["required_sections"].items()]}
        write_json(output / f"{mode}.json", draft)
    # Explicit fixture, not a customer template and not evidence of customer template fidelity.
    template = default_document()
    template.sections[0].left_margin = Mm(26)
    template.sections[0].right_margin = Mm(22)
    template.styles["Heading 1"].font.size = Pt(14)
    template.sections[0].header.paragraphs[0].text = "技术资料"
    template.add_paragraph("{{ title }}", "Title")
    template.add_paragraph("{{p body }}")
    template.save(output / "reference-fixture.docx")
    write_json(output / "fixture-info.json", {"template": "Created for testing, not user-provided", "project": str(project), "review_scope": "source analysis, no performance claim", "sources": paths})


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--project", required=True)
    p.add_argument("--out", required=True)
    a = p.parse_args()
    prepare(a.project, a.out)
